# -*- coding: utf-8 -*-
"""スライド14（移乗の4つの方法）を作り直す。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY  = RGBColor(0x1A, 0x52, 0x76)
GREEN = RGBColor(0x1D, 0x9E, 0x75)
DARK  = RGBColor(0x1C, 0x28, 0x33)
GRAY  = RGBColor(0x56, 0x65, 0x73)
LITE  = RGBColor(0xF4, 0xF8, 0xFA)
GRNLT = RGBColor(0xEA, 0xF7, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
F = "Meiryo"

prs = Presentation("houmon-nyuyoku-slides.pptx")
sld = prs.slides[13]
for sh in list(sld.shapes):
    sh._element.getparent().remove(sh._element)

def box(x, y, w, h, fill=None, line=None, radius=0.10):
    sp = sld.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Inches(x), Inches(y), Inches(w), Inches(h))
    sp.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def text(x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sld.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (s, size, color, bold, space) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.15
        if space:
            p.space_before = Pt(space)
        r = p.add_run(); r.text = s
        r.font.name = F; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

# ---------- ヘッダー ----------
box(0.62, 0.36, 1.62, 0.34, fill=GREEN, radius=0.05)
text(0.62, 0.36, 1.62, 0.34, [("DEMO", 12, WHITE, True, 0)],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(2.40, 0.36, 6.00, 0.34, [("移乗の実演", 13, GRAY, False, 0)], anchor=MSO_ANCHOR.MIDDLE)
text(0.62, 0.78, 12.09, 0.60, [("お体を、どうお運びするか", 32, NAVY, True, 0)])
text(0.62, 1.52, 12.09, 0.32,
     [("浴槽へお移しする方法は4つ。いずれも3名で行い、その方の状態に合わせて選びます。",
       13.5, GRAY, False, 0)])

# ---------- カード ----------
CARDS = [
    ("抱き上げ移動", "体格が小さく、拘縮の少ない方",
     "選ばないのは", "体重の重い方", "いちばん短時間"),
    ("担架移動", "体の大きい方、気管切開の方、全介助の方",
     "必ずこの方法", "人工呼吸器を外せない方。回路が引っ張られないように", "いちばん安全"),
    ("上下移動", "車椅子で生活されている方",
     "支え方", "上半身・臀部・下肢を3名で支え、高さを合わせます", "車椅子から直接"),
    ("座位回転", "介助で立位がとれ、少し歩ける方",
     "選ばないのは", "立位がとれない方", "ご本人の力を残す"),
]
W, GAP, X0, Y, H = 2.95, 0.10, 0.62, 2.05, 3.45
for i, (name, who, lab2, note, merit) in enumerate(CARDS):
    x = X0 + i * (W + GAP)
    box(x, Y, W, H, fill=LITE)
    box(x, Y, W, 0.58, fill=NAVY)
    text(x, Y, W, 0.58, [(name, 17, WHITE, True, 0)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(x + 0.20, Y + 0.72, W - 0.40, 0.24, [("こんな方に", 10.5, GREEN, True, 0)])
    text(x + 0.20, Y + 1.00, W - 0.40, 0.78, [(who, 12.5, DARK, False, 0)])
    text(x + 0.20, Y + 1.86, W - 0.40, 0.24, [(lab2, 10.5, NAVY, True, 0)])
    text(x + 0.20, Y + 2.14, W - 0.40, 0.78, [(note, 12, GRAY, False, 0)])
    box(x + 0.20, Y + 2.92, W - 0.40, 0.44, fill=GRNLT)
    text(x + 0.20, Y + 2.92, W - 0.40, 0.44, [(merit, 12.5, GREEN, True, 0)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ---------- 下の帯 ----------
box(0.62, 5.72, 12.09, 1.00, fill=GREEN)
text(0.92, 5.72, 11.49, 1.00,
     [("方法は固定ではありません。", 17, WHITE, True, 0),
      ("事前訪問で決め、やってみてより安全・安楽な方へ変えます。ADLが上がれば座位回転へ。"
       "変更のつど、ご本人とご家族に確認しています。", 13.5, WHITE, False, 5)],
     anchor=MSO_ANCHOR.MIDDLE)

prs.save("houmon-nyuyoku-slides.pptx")
E = 914400.0
s = Presentation("houmon-nyuyoku-slides.pptx").slides[13]
print("図形数:", len(s.shapes))
print("最下端 %.2f in ／ 右端 %.2f in" %
      (max((sh.top + sh.height) / E for sh in s.shapes),
       max((sh.left + sh.width) / E for sh in s.shapes)))
