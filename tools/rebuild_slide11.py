# -*- coding: utf-8 -*-
"""スライド11（在宅チームの一員として）を作り直す。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import copy

NAVY   = RGBColor(0x1A, 0x52, 0x76)
GREEN  = RGBColor(0x1D, 0x9E, 0x75)
DARK   = RGBColor(0x1C, 0x28, 0x33)
GRAY   = RGBColor(0x56, 0x65, 0x73)
BGLITE = RGBColor(0xF4, 0xF8, 0xFA)
GRN_LT = RGBColor(0xEA, 0xF7, 0xF1)
GRN_MD = RGBColor(0xD6, 0xF3, 0xE6)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
FONT   = "Meiryo"

prs = Presentation("houmon-nyuyoku-slides.pptx")
sld = prs.slides[10]

# ヘッダー（CHAPTER 4 / 見出し）だけ残して、以下を作り直す
for sh in list(sld.shapes)[4:]:
    sh._element.getparent().remove(sh._element)

def box(x, y, w, h, fill=None, line=None, radius=0.10):
    sp = sld.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def text(x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines = [(文字, サイズpt, 色, 太字, 行間前pt), ...]"""
    tb = sld.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (s, size, color, bold, space) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.space_before = Pt(space)
        r = p.add_run(); r.text = s
        r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

# ---------- 上段：連携図 ----------
box(0.62, 1.80, 2.60, 1.70, fill=GREEN)
text(0.72, 1.80, 2.40, 1.70,
     [("訪問入浴介護", 15, WHITE, True, 0),
      ("看護師1名＋介護2名", 10.5, WHITE, False, 6)],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

PARTNERS = [
    (3.90, 1.80, "主治医・在宅医",        "創部・全身状態の報告"),
    (8.40, 1.80, "訪問看護ステーション",  "観察所見の共有"),
    (3.90, 2.72, "ケアマネジャー",        "状態変化の報告"),
    (8.40, 2.72, "デイサービス・訪問介護", "入浴日の分担"),
]
for x, y, name, flow in PARTNERS:
    box(x, y, 4.30, 0.78, fill=BGLITE)
    text(x + 0.22, y, 2.30, 0.78, [(name, 13.5, NAVY, True, 0)], anchor=MSO_ANCHOR.MIDDLE)
    text(x + 2.45, y, 1.70, 0.78, [(flow, 10.5, GRAY, False, 0)],
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    if x < 5:
        text(x - 0.52, y, 0.44, 0.78, [("⇄", 17, GREEN, True, 0)],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ---------- 下段：デイサービスとの関係（ADLの時間軸）----------
text(0.62, 3.86, 12.08, 0.34,
     [("デイサービスとは「競合」ではなく、引き継ぎと併用です", 14, NAVY, True, 0)])

# 併用期の帯（いちばん後ろに置く）
box(5.30, 4.46, 1.42, 1.54, fill=GRN_MD)
text(5.30, 4.16, 1.42, 0.28, [("併用期", 10.5, GREEN, True, 0)], align=PP_ALIGN.CENTER)

box(0.62, 4.60, 6.10, 0.62, fill=BGLITE, line=NAVY)
text(0.92, 4.60, 5.50, 0.62, [("デイサービス（通所）での入浴", 13, NAVY, True, 0)],
     anchor=MSO_ANCHOR.MIDDLE)

box(5.30, 5.38, 7.40, 0.62, fill=GREEN)
text(5.60, 5.38, 6.80, 0.62, [("訪問入浴（自宅での全身浴）", 13, WHITE, True, 0)],
     anchor=MSO_ANCHOR.MIDDLE)

# 時間の矢印
ar = sld.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(0.62), Inches(6.16), Inches(12.08), Inches(0.14))
ar.fill.solid(); ar.fill.fore_color.rgb = RGBColor(0xC7, 0xD2, 0xDA)
ar.line.fill.background(); ar.shadow.inherit = False

text(0.62, 6.40, 5.00, 0.30, [("ADLが保たれ、通所できる時期", 10.5, GRAY, False, 0)])
text(7.70, 6.40, 5.00, 0.30, [("ADLが低下し、通所が難しくなる時期", 10.5, GRAY, False, 0)],
     align=PP_ALIGN.RIGHT)

text(0.62, 6.82, 12.08, 0.40,
     [("役割を分け合うことで、ご自宅で過ごせる期間が延びます。", 13.5, GREEN, True, 0)])

prs.save("houmon-nyuyoku-slides.pptx")
print("slide 11 を作り直しました")
