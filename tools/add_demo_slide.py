# -*- coding: utf-8 -*-
"""移乗の4つの方法のスライドを、13枚目のうしろに差し込む。"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY  = RGBColor(0x1A, 0x52, 0x76)
GREEN = RGBColor(0x1D, 0x9E, 0x75)
DARK  = RGBColor(0x1C, 0x28, 0x33)
GRAY  = RGBColor(0x56, 0x65, 0x73)
LITE  = RGBColor(0xF4, 0xF8, 0xFA)
GRNLT = RGBColor(0xEA, 0xF7, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
F = "Meiryo"

prs = Presentation("houmon-nyuyoku-slides.pptx")
sld = prs.slides.add_slide(prs.slides[12].slide_layout)

def box(x, y, w, h, fill=None, line=None, radius=0.10):
    sp = sld.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Inches(x), Inches(y), Inches(w), Inches(h))
    sp.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def text(x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sld.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (s, size, color, bold, space) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.space_before = Pt(space)
        r = p.add_run(); r.text = s
        r.font.name = F; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

# ヘッダー
box(0.62, 0.36, 1.62, 0.34, fill=GREEN, radius=0.05)
text(0.62, 0.36, 1.62, 0.34, [("DEMO", 12, WHITE, True, 0)],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(2.40, 0.36, 6.00, 0.34, [("移乗の実演", 13, GRAY, False, 0)], anchor=MSO_ANCHOR.MIDDLE)
text(0.62, 0.82, 12.09, 0.62, [("お体を、どうお運びするか", 34, NAVY, True, 0)])
text(0.62, 1.60, 12.09, 0.34,
     [("浴槽へお移しする方法は4つ。その方の状態に合わせて選びます。", 14, GRAY, False, 0)])

METHODS = [
    ("1", "抱き上げ移動", "体格が小さく、拘縮の\n少ない方",      "いちばん短時間"),
    ("2", "担架移動",     "体の大きい方、\n全介助の方",          "いちばん安全"),
    ("3", "上下移動",     "ベッドと浴槽の高さを\n合わせ、水平に",  "体への負担が少ない"),
    ("4", "座位回転",     "座位が保てる方",                      "ご本人の力を残す"),
]
W, GAP, X0, Y, H = 2.95, 0.10, 0.62, 2.20, 3.05
for i, (num, name, who, merit) in enumerate(METHODS):
    x = X0 + i * (W + GAP)
    box(x, Y, W, H, fill=LITE)
    box(x, Y, W, 0.62, fill=NAVY)
    text(x, Y, W, 0.62, [(name, 17, WHITE, True, 0)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(x + 0.22, Y + 0.80, W - 0.44, 0.26,
         [("こんな方に", 11, GREEN, True, 0)])
    text(x + 0.22, Y + 1.12, W - 0.44, 1.10,
         [(l, 13.5, DARK, False, 0 if j == 0 else 4) for j, l in enumerate(who.split("\n"))])
    box(x + 0.22, Y + 2.28, W - 0.44, 0.50, fill=GRNLT)
    text(x + 0.22, Y + 2.28, W - 0.44, 0.50, [(merit, 13, GREEN, True, 0)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

box(0.62, 5.60, 12.09, 0.72, fill=GREEN)
text(0.62, 5.60, 12.09, 0.72,
     [("どの方法でお運びするかを、お一人おひとり決めています。安全に、楽に。", 18, WHITE, True, 0)],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# 新しいスライドのノート欄には本文の枠が無いので、既存のノートから複製する
NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
dst = sld.notes_slide
if dst.notes_text_frame is None:
    src = prs.slides[12].notes_slide
    for sh in src.shapes:
        ph = sh.element.find(".//" + NS + "ph")
        if ph is not None and ph.get("type") == "body":
            dst.shapes._spTree.append(copy.deepcopy(sh.element))
            break
dst.notes_text_frame.text = """［42:30〜50:30／8分］★移乗デモンストレーション
・冨永＝説明と進行／高崎・武末＝実演
・導入「ここで少し、体を動かしていただきます」
・4つを実演しながら説明。名前だけでなく「どんな方に使うか」を必ず添える
  抱き上げ＝小柄・拘縮少ない／担架＝大柄・全介助／
  上下＝高さを合わせ水平に／座位回転＝座位が保てる方
・体験（3分）「どなたか、お体を預けてみていただけませんか」
  → 挙手を待つ。指名しない。腰・首の既往を確認。1人だけ。担架がおすすめ
  → 手が挙がらなければ「ではスタッフでお見せします」に即切り替え
・締め「持ち上げられる側は、怖いものです。だからその方ごとに決めています」
【押していたら】体験を省いて実演だけ −3分／2方法に絞る さらに−2分"""

# 13枚目のうしろ（14枚目）へ移動する
ids = prs.slides._sldIdLst
ids.insert(13, ids[-1])
prs.save("houmon-nyuyoku-slides.pptx")

p2 = Presentation("houmon-nyuyoku-slides.pptx")
print("スライド数:", len(p2.slides.__iter__.__self__._sldIdLst))
for i, s in enumerate(p2.slides, 1):
    t = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    print("%2d  %s" % (i, " / ".join(t)[:46].replace("\n", " ")))
